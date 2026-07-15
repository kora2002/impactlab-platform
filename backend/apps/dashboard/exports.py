import io
from django.http import HttpResponse
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from django.db.models import Count, Sum, Q
from apps.beneficiaires.models import Beneficiaire
from apps.programmes.models import Programme
from apps.inscriptions.models import Inscription, Progression
from apps.suivi.models import SuiviInsertion
from apps.financements.models import Financement


def get_indicateurs():
    """Calcule tous les indicateurs du dashboard"""
    total_beneficiaires = Beneficiaire.objects.count()
    nombre_femmes = Beneficiaire.objects.filter(genre='femme').count()
    nombre_hommes = Beneficiaire.objects.filter(genre='homme').count()
    taux_parite = round((nombre_femmes / total_beneficiaires) * 100, 2) if total_beneficiaires > 0 else 0

    total_programmes = Programme.objects.count()
    programmes_actifs = Programme.objects.filter(statut='actif').count()

    total_inscriptions = Inscription.objects.count()
    inscriptions_validees = Inscription.objects.filter(statut='validee').count()
    total_abandons = Progression.objects.filter(statut='abandonne').count()
    taux_abandon = round((total_abandons / total_inscriptions) * 100, 2) if total_inscriptions > 0 else 0

    total_inseres = SuiviInsertion.objects.exclude(
        statut_insertion__in=['non_repondu', 'sans_solution']
    ).values('beneficiaire').distinct().count()
    taux_insertion = round((total_inseres / total_beneficiaires) * 100, 2) if total_beneficiaires > 0 else 0

    total_prevu = Financement.objects.aggregate(t=Sum('montant_prevu'))['t'] or 0
    total_realise = Financement.objects.aggregate(t=Sum('montant_realise'))['t'] or 0

    repartition_localite = Beneficiaire.objects.values('localite').annotate(
        total=Count('id'),
        femmes=Count('id', filter=Q(genre='femme'))
    ).order_by('-total')

    return {
        'total_beneficiaires': total_beneficiaires,
        'nombre_femmes': nombre_femmes,
        'nombre_hommes': nombre_hommes,
        'taux_parite': taux_parite,
        'total_programmes': total_programmes,
        'programmes_actifs': programmes_actifs,
        'total_inscriptions': total_inscriptions,
        'inscriptions_validees': inscriptions_validees,
        'taux_abandon': taux_abandon,
        'total_inseres': total_inseres,
        'taux_insertion': taux_insertion,
        'total_prevu': total_prevu,
        'total_realise': total_realise,
        'repartition_localite': list(repartition_localite),
    }


def export_excel(request):
    """Génère un fichier Excel du rapport d'impact"""
    ind = get_indicateurs()

    wb = Workbook()

    # ── Feuille 1 : Indicateurs globaux ──
    ws1 = wb.active
    ws1.title = "Indicateurs globaux"

    # Style entête
    header_fill = PatternFill(start_color="1F4E5F", end_color="1F4E5F", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    center = Alignment(horizontal="center", vertical="center")

    # Titre
    ws1.merge_cells("A1:C1")
    ws1["A1"] = "RAPPORT D'IMPACT — IMPACT'LAB GDC"
    ws1["A1"].font = Font(bold=True, size=14, color="1F4E5F")
    ws1["A1"].alignment = center
    ws1.row_dimensions[1].height = 30

    ws1.append([])

    # Headers
    headers = ["Indicateur", "Valeur", "Détail"]
    ws1.append(headers)
    for col, header in enumerate(headers, 1):
        cell = ws1.cell(row=3, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    # Données
    donnees = [
        ["Total bénéficiaires", ind['total_beneficiaires'], ""],
        ["Nombre de femmes", ind['nombre_femmes'], f"{ind['taux_parite']}% de femmes"],
        ["Nombre d'hommes", ind['nombre_hommes'], ""],
        ["Taux de parité", f"{ind['taux_parite']}%", "Objectif : 43%"],
        ["Programmes actifs", ind['programmes_actifs'], f"sur {ind['total_programmes']} programmes"],
        ["Total inscriptions", ind['total_inscriptions'], ""],
        ["Inscriptions validées", ind['inscriptions_validees'], ""],
        ["Taux d'abandon", f"{ind['taux_abandon']}%", ""],
        ["Bénéficiaires insérés", ind['total_inseres'], f"Taux : {ind['taux_insertion']}%"],
        ["Financements prévus", f"{ind['total_prevu']:,.0f} FCFA", ""],
        ["Financements réalisés", f"{ind['total_realise']:,.0f} FCFA", ""],
    ]

    for row_data in donnees:
        ws1.append(row_data)

    # Largeur colonnes
    ws1.column_dimensions["A"].width = 30
    ws1.column_dimensions["B"].width = 20
    ws1.column_dimensions["C"].width = 25

    # ── Feuille 2 : Répartition géographique ──
    ws2 = wb.create_sheet("Répartition géographique")
    ws2.append(["Localité", "Total bénéficiaires", "Femmes"])
    for col in range(1, 4):
        cell = ws2.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    for loc in ind['repartition_localite']:
        ws2.append([loc['localite'] or "Non renseigné", loc['total'], loc['femmes']])

    ws2.column_dimensions["A"].width = 25
    ws2.column_dimensions["B"].width = 20
    ws2.column_dimensions["C"].width = 15

    # ── Feuille 3 : Bénéficiaires ──
    ws3 = wb.create_sheet("Bénéficiaires")
    headers_benef = ["Nom", "Prénom", "Genre", "Localité", "Téléphone", "Statut pro"]
    ws3.append(headers_benef)
    for col, h in enumerate(headers_benef, 1):
        cell = ws3.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font

    for b in Beneficiaire.objects.all():
        ws3.append([b.nom, b.prenom, b.get_genre_display(), b.localite, b.telephone, b.statut_pro])

    for col in ["A", "B", "C", "D", "E", "F"]:
        ws3.column_dimensions[col].width = 18

    # Réponse HTTP
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="rapport_impact_impactlab.xlsx"'
    wb.save(response)
    return response


def export_pdf(request):
    """Génère un fichier PDF du rapport d'impact"""
    ind = get_indicateurs()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    elements = []

    # Titre
    title_style = styles["Title"]
    title_style.textColor = colors.HexColor("#1F4E5F")
    elements.append(Paragraph("RAPPORT D'IMPACT", title_style))
    elements.append(Paragraph("Impact'Lab GDC", styles["Heading2"]))
    elements.append(Spacer(1, 20))

    # Tableau indicateurs globaux
    elements.append(Paragraph("Indicateurs globaux", styles["Heading3"]))
    elements.append(Spacer(1, 10))

    data = [
        ["Indicateur", "Valeur"],
        ["Total bénéficiaires", str(ind['total_beneficiaires'])],
        ["Femmes", f"{ind['nombre_femmes']} ({ind['taux_parite']}%)"],
        ["Hommes", str(ind['nombre_hommes'])],
        ["Taux de parité", f"{ind['taux_parite']}% (objectif 43%)"],
        ["Programmes actifs", f"{ind['programmes_actifs']} / {ind['total_programmes']}"],
        ["Total inscriptions", str(ind['total_inscriptions'])],
        ["Inscriptions validées", str(ind['inscriptions_validees'])],
        ["Taux d'abandon", f"{ind['taux_abandon']}%"],
        ["Bénéficiaires insérés", f"{ind['total_inseres']} ({ind['taux_insertion']}%)"],
        ["Financements prévus", f"{ind['total_prevu']:,.0f} FCFA"],
        ["Financements réalisés", f"{ind['total_realise']:,.0f} FCFA"],
    ]

    table = Table(data, colWidths=[300, 180])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 11),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F7FA")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
        ("PADDING", (0, 0), (-1, -1), 8),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 20))

    # Tableau répartition géographique
    elements.append(Paragraph("Répartition géographique", styles["Heading3"]))
    elements.append(Spacer(1, 10))

    geo_data = [["Localité", "Total", "Femmes"]]
    for loc in ind['repartition_localite']:
        geo_data.append([loc['localite'] or "Non renseigné", str(loc['total']), str(loc['femmes'])])

    geo_table = Table(geo_data, colWidths=[250, 115, 115])
    geo_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F7FA")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
        ("PADDING", (0, 0), (-1, -1), 8),
    ]))
    elements.append(geo_table)

    doc.build(elements)
    buffer.seek(0)

    response = HttpResponse(buffer, content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="rapport_impact_impactlab.pdf"'
    return response