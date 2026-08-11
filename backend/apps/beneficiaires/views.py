from rest_framework import generics, permissions
from rest_framework.views import APIView
from rest_framework.response import Response
from django.db.models import Q
from .models import Beneficiaire, Structure, MembreStructure
from .serializers import (
    BeneficiaireSerializer,
    BeneficiaireListSerializer,
    BeneficiaireHistoriqueSerializer,
    StructureSerializer,
    StructureListSerializer,
    MembreStructureSerializer,
)
from .connectors import importer_organisations_asso_pro
from .connectors import importer_organisations_asso_pro, importer_porteurs_sageo

class BeneficiaireListCreateView(generics.ListCreateAPIView):
    """
    GET  /api/beneficiaires/   → liste (avec recherche)
    POST /api/beneficiaires/   → créer une fiche
    """
    permission_classes = [permissions.IsAuthenticated]
    queryset = Beneficiaire.objects.all().order_by('nom')

    def get_serializer_class(self):
        if self.request.method == 'GET':
            return BeneficiaireListSerializer
        return BeneficiaireSerializer

    def get_queryset(self):
        qs = Beneficiaire.objects.all().order_by('nom')
        recherche = self.request.query_params.get('recherche')
        if recherche:
            qs = qs.filter(
                Q(nom__icontains=recherche) |
                Q(prenom__icontains=recherche) |
                Q(telephone__icontains=recherche) |
                Q(localite__icontains=recherche)
            )
        genre = self.request.query_params.get('genre')
        if genre:
            qs = qs.filter(genre=genre)
        return qs


class BeneficiaireDetailView(generics.RetrieveUpdateDestroyAPIView):
    """GET/PUT/DELETE /api/beneficiaires/<id>/"""
    serializer_class   = BeneficiaireSerializer
    permission_classes = [permissions.IsAuthenticated]
    queryset           = Beneficiaire.objects.all()


class BeneficiaireHistoriqueView(generics.RetrieveAPIView):
    """GET /api/beneficiaires/<id>/historique/"""
    serializer_class   = BeneficiaireHistoriqueSerializer
    permission_classes = [permissions.IsAuthenticated]
    queryset           = Beneficiaire.objects.all()


class StructureListCreateView(generics.ListCreateAPIView):
    """
    GET  /api/structures/   → liste les structures
    POST /api/structures/   → créer une structure
    """
    permission_classes = [permissions.IsAuthenticated]
    queryset = Structure.objects.all().order_by('nom')

    def get_serializer_class(self):
        if self.request.method == 'GET':
            return StructureListSerializer
        return StructureSerializer

    def get_queryset(self):
        qs = Structure.objects.all().order_by('nom')
        type_structure = self.request.query_params.get('type')
        if type_structure:
            qs = qs.filter(type=type_structure)
        return qs


class StructureDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class   = StructureSerializer
    permission_classes = [permissions.IsAuthenticated]
    queryset           = Structure.objects.all()


class AjouterMembreStructureView(generics.CreateAPIView):
    """POST /api/structures/<structure_id>/membres/"""
    serializer_class   = MembreStructureSerializer
    permission_classes = [permissions.IsAuthenticated]

    def perform_create(self, serializer):
        structure = Structure.objects.get(id=self.kwargs['structure_id'])
        serializer.save(structure=structure)


class ImportExcelBeneficiairesView(APIView):
    """POST /api/beneficiaires/import-excel/"""
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        import openpyxl
        from rest_framework import status

        fichier = request.FILES.get('fichier')
        if not fichier:
            return Response(
                {"detail": "Aucun fichier fourni."},
                status=status.HTTP_400_BAD_REQUEST
            )
        try:
            wb = openpyxl.load_workbook(fichier)
            ws = wb.active
        except Exception:
            return Response(
                {"detail": "Fichier Excel invalide."},
                status=status.HTTP_400_BAD_REQUEST
            )

        crees   = 0
        erreurs = []

        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                nom       = str(row[0]).strip() if row[0] else None
                prenom    = str(row[1]).strip() if row[1] else None
                genre     = str(row[2]).strip().lower() if row[2] else None
                telephone = str(row[3]).strip() if row[3] else None
                email     = str(row[4]).strip() if row[4] else ""
                localite  = str(row[5]).strip() if row[5] else ""

                if not nom or not prenom or not telephone:
                    erreurs.append(f"Ligne {i} : nom, prénom et téléphone obligatoires.")
                    continue

                if genre not in ['homme', 'femme']:
                    erreurs.append(f"Ligne {i} : genre invalide (doit être 'homme' ou 'femme').")
                    continue

                if Beneficiaire.objects.filter(telephone=telephone).exists():
                    erreurs.append(f"Ligne {i} : bénéficiaire avec téléphone {telephone} déjà existant.")
                    continue

                Beneficiaire.objects.create(
                    nom=nom, prenom=prenom, genre=genre,
                    telephone=telephone, email=email, localite=localite,
                )
                crees += 1

            except Exception as e:
                erreurs.append(f"Ligne {i} : erreur — {str(e)}")

        return Response({
            "crees": crees,
            "erreurs": erreurs,
            "message": f"{crees} bénéficiaire(s) importé(s) avec succès."
        })


class TelechargerModeleExcelView(APIView):
    """GET /api/beneficiaires/modele-excel/"""
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        import openpyxl
        from django.http import HttpResponse
        from openpyxl.styles import Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Bénéficiaires"

        headers = ["nom *", "prenom *", "genre * (homme/femme)", "telephone *",
                   "email", "localite", "quartier", "date_naissance (AAAA-MM-JJ)",
                   "nationalite", "niveau_etudes", "statut_pro",
                   "contact_urgence_nom", "contact_urgence_tel"]
        ws.append(headers)

        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="1F4E5F", end_color="1F4E5F", fill_type="solid")
            ws.column_dimensions[cell.column_letter].width = 20

        ws.append(["Koné", "Fatou", "femme", "0700000001", "fatou@email.com",
                   "Abidjan", "Yopougon", "2000-05-15", "Ivoirienne",
                   "Licence", "demandeur_emploi", "Koné Mamadou", "0700000002"])

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="modele_beneficiaires.xlsx"'
        wb.save(response)
        return response


class ImportAssoproView(APIView):
    """POST /api/beneficiaires/import-asso-pro/"""
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        resultat = importer_organisations_asso_pro()
        return Response(resultat)
    

class ImportSageoView(APIView):
    """POST /api/beneficiaires/import-sageo/"""
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        resultat = importer_porteurs_sageo()
        return Response(resultat)